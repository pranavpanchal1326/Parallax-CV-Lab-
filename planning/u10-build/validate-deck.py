"""Writes submission/presentation-validation.json for the editorial deck:
ZIP CRC, reopen, slide/notes counts, text and notes parity with the source JSON, alt text."""
import json, hashlib, zipfile
from pathlib import Path
from pptx import Presentation
root = Path(__file__).resolve().parents[2]
deck = root / 'submission/PARALLAX-Presentation-Final.pptx'
with zipfile.ZipFile(deck) as z:
    assert z.testzip() is None
    parts = z.namelist()
prs = Presentation(str(deck))
content = json.loads((root / 'submission/source/presentation-content.json').read_text(encoding='utf-8'))
fonts, alts = set(), 0
assert len(prs.slides) == len(content) == 15
for i, sl in enumerate(prs.slides):
    text = ' '.join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame)
    for sh in sl.shapes:
        if sh.has_text_frame:
            for par in sh.text_frame.paragraphs:
                fonts.update(r.font.name for r in par.runs)
        if sh.shape_type == 13 and sh._element.nvPicPr.cNvPr.get('descr'):
            alts += 1
    assert content[i]['title'] in text.replace('\n', ''), i
    assert all(b in text for b in content[i]['body']), i
    assert content[i]['notes'] in sl.notes_slide.notes_text_frame.text, i
slides = sorted((root / 'submission/deck-slides').glob('slide-*.png'))
assert len(slides) == 15 and all(p.stat().st_mtime >= deck.stat().st_mtime - 1 for p in slides), 'slide images older than the deck'
rec = {'schemaVersion': 'presentation-finalization.v2', 'revision': 'U10-R3', 'date': '2026-10-05',
       'finalSha256': hashlib.sha256(deck.read_bytes()).hexdigest(), 'byteCount': deck.stat().st_size,
       'packageIntegrity': {'status': 'pass', 'slide_count': len(prs.slides), 'zipCrc': 'pass', 'parts': len(parts), 'notesSlides': 15, 'picturesWithAltText': alts,
                            'claim_boundary': 'ZIP CRC, python-pptx reopen, slide/notes counts and text parity with presentation-content.json.'},
       'contentParity': {'titlesAndBodyLinesMatchSource': True, 'speakerNotesMatchSource': True, 'plannedSeconds': sum(r['seconds'] for r in content)},
       'fontPolicy': {'observed_font_families': sorted(fonts), 'note': 'Georgia, Arial and Consolas ship with Windows and Office; another platform may substitute Consolas.'},
       'nativePowerPoint': {'opened': True, 'version': '16.0', 'slidesExported': 15, 'method': 'planning/u10-build/render-deck.ps1 (COM, read-only, no window); submission/deck-slides holds these exports',
                            'scope': 'One machine. Slide show playback and other PowerPoint versions were not exercised.'},
       'previousRevision': {'preservedAt': 'planning/revisions/U10-R1-before-editorial-20261005/'},
       'scope': 'Editorial layout; wording, order, timings and notes unchanged from U10-R1. Interface figures re-captured after the product polish.'}
(root / 'submission/presentation-validation.json').write_text(json.dumps(rec, indent=2), encoding='utf-8', newline='\n')
print(rec['finalSha256'][:16], sorted(fonts), alts)
