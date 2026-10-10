"""U10-R2 editorial deck. Same fifteen slides, wording, notes and timings as
submission/source/presentation-content.json; only the layout language changes
(paper, ink, hairline rules, serif headlines, one ultramarine spot colour)."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'submission'
content = json.loads((OUT / 'source/presentation-content.json').read_text(encoding='utf-8'))
assert len(content) == 15 and sum(r['seconds'] for r in content) == 900
team_file = OUT / 'source/team.json'
team = json.loads(team_file.read_text(encoding='utf-8')) if team_file.exists() else None

PAPER, CARD, INK, INK2, SPOT, RULE = (RGBColor.from_string(x) for x in
                                      ('F4F1E9', 'FBFAF5', '20232B', '535B66', '3346D3', 'B9B6AE'))
SERIF, SANS, MONO = 'Georgia', 'Arial', 'Consolas'
W, H, M = 13.333, 7.5, 0.7

# Presentation-only: which closing phrase of each unchanged title is set in italic spot colour.
EMPHASIS = [None, 'a free station', 'Fourteen sources', 'ownership visible', 'inspectable',
            'seventy square metres', 'can support', 'assumptions and owners',
            'ambiguous sampling choice', 'bounded comparison', 'until verification',
            'with its revision', 'complete dependencies', 'original behavior', 'the whole submission']
SECTION = ['Proposal', 'Problem', 'Research', 'Journey', 'Wireframes', 'Space', 'Dashboard',
           'Sustainability', 'Demonstration · E01', 'Demonstration · E06', 'Recovery', 'Local AI',
           'Operations', 'Verification', 'Handoff']
CAPTION = {
    1: 'Conceptual P4-D render. Not a photograph of an installed room.',
    3: 'Secondary-research page, desktop capture. Zero local participants.',
    4: 'Student-eye render of the conceptual P4-D model.',
    5: 'WF03. Complete-kit booking and conflict recovery.',
    6: 'P4-D orientation plan. Assumed 10 × 7 m envelope.',
    7: 'Current dashboard. Sensor connections are explicitly zero.',
    8: 'Ventilation plan layer. Conceptual proposal, no site survey.',
    9: 'E01 rehearsal at 3 Hz. Generated teaching fixture, not a measurement.',
    10: 'E06 authored lighting comparison. Illustration, not detector output.',
    11: 'Fault sandbox, stage 8. Early release is refused.',
    12: 'Real local Ollama answer with its guide citation.',
    13: 'Independent operations exercise. Local allocation only.',
    14: 'Typed equipment lookup at phone width.',
    15: 'The current submission entry.',
}
EXTRA_IMAGE = {15: 'submission/assets/submission-entry-top.png'}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'PARALLAX - Signal & Computer Vision Lab'
prs.core_properties.author = 'PARALLAX project workspace'
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def run(par, text, font, size, color=INK, bold=False, italic=False, spacing=None):
    r = par.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    if spacing is not None:
        r._r.get_or_add_rPr().set('spc', str(spacing))
    return r


def line(slide, x1, y1, x2, y2, color=INK, weight=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    c._element.spPr.append(c._element.spPr.makeelement(qn('a:effectLst'), {}))  # no theme shadow on rules
    return c


def rect(slide, x, y, w, h, fill=None, outline=None, weight=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if outline is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = outline
        s.line.width = Pt(weight)
    return s


def headline(slide, title, emphasis, x, y, w, h, size):
    tf = box(slide, x, y, w, h)
    par = tf.paragraphs[0]
    par.line_spacing = 1.04
    if emphasis and emphasis in title:
        before, after = title.split(emphasis, 1)
        if before:
            run(par, before, SERIF, size)
        run(par, emphasis, SERIF, size, SPOT, italic=True)
        if after:
            run(par, after, SERIF, size)
    else:
        run(par, title, SERIF, size)


def figure(slide, rel, x, y, w, h, number, caption, alt):
    """Hairline plate with crop marks, the picture fitted inside, and a numbered caption."""
    rect(slide, x, y, w, h, fill=CARD, outline=RULE, weight=0.5)
    pad = 0.12
    with PILImage.open(ROOT / rel) as im:
        iw, ih = im.size
    scale = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
    pw, ph = iw * scale, ih * scale
    pic = slide.shapes.add_picture(str(ROOT / rel), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                   Inches(pw), Inches(ph))
    pic._element.nvPicPr.cNvPr.set('descr', alt)
    t = 0.16
    for cx, cy, dx, dy in ((x, y, -1, -1), (x + w, y, 1, -1), (x, y + h, -1, 1), (x + w, y + h, 1, 1)):
        line(slide, cx + dx * 0.05, cy, cx + dx * (0.05 + t), cy, INK, 0.5)
        line(slide, cx, cy + dy * 0.05, cx, cy + dy * (0.05 + t), INK, 0.5)
    tf = box(slide, x, y + h + 0.3, w, 0.3)
    par = tf.paragraphs[0]
    run(par, f'FIG. {number:02d}', MONO, 9, SPOT, spacing=120)
    run(par, '   ' + caption, SANS, 10, INK2)


def signal_rule(slide, x, y, w, period=1.5, amp=0.11):
    """The sampled 2 Hz reference as a decorative rule: eight spot-colour samples per period."""
    import math
    step = period / 8
    k = step / 2
    while k < w:
        d = 0.06
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + k - d / 2), Inches(y - amp * math.sin(2 * math.pi * k / period) - d / 2), Inches(d), Inches(d))
        dot.shadow.inherit = False
        dot.fill.solid()
        dot.fill.fore_color.rgb = SPOT
        dot.line.fill.background()
        k += step


def chrome(slide, i, start, end):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    rect(slide, M, 0.40, 0.13, 0.13, fill=SPOT)
    par = box(slide, M + 0.24, 0.36, 4, 0.25).paragraphs[0]
    run(par, 'PARALLAX', SANS, 10.5, INK, bold=True, spacing=300)
    par = box(slide, W - M - 6, 0.37, 6, 0.25).paragraphs[0]
    par.alignment = PP_ALIGN.RIGHT
    run(par, f'{SECTION[i].upper()}   ·   {i + 1:02d} / 15', MONO, 9, INK2, spacing=120)
    line(slide, M, 0.72, W - M, 0.72, INK, 0.75)
    line(slide, M, 6.92, W - M, 6.92, RULE, 0.5)
    par = box(slide, M, 7.0, 7, 0.22).paragraphs[0]
    run(par, 'SIGNAL & COMPUTER VISION LAB   ·   70 m²   ·   LOCAL PROPOSAL', MONO, 8, INK2, spacing=100)
    par = box(slide, W - M - 4, 7.0, 4, 0.22).paragraphs[0]
    par.alignment = PP_ALIGN.RIGHT
    stamp = lambda s: f'{s // 60}:{s % 60:02d}'
    run(par, f'{stamp(start)} – {stamp(end)}', MONO, 8, INK2, spacing=100)


def points(slide, body, x, y, w, step, size=15):
    for j, text in enumerate(body):
        top = y + j * step
        line(slide, x, top, x + w, top, RULE, 0.5)
        par = box(slide, x, top + 0.13, 0.5, 0.3).paragraphs[0]
        run(par, f'{j + 1:02d}', MONO, 10, SPOT, spacing=80)
        par = box(slide, x + 0.55, top + 0.1, w - 0.55, step - 0.15).paragraphs[0]
        par.line_spacing = 1.18
        run(par, text, SANS, size)


minute = 0
for i, row in enumerate(content):
    slide = prs.slides.add_slide(blank)
    start, minute = minute, minute + row['seconds']
    chrome(slide, i, start, minute)
    n = i + 1
    image = row.get('image') or EXTRA_IMAGE.get(n)
    if n == 1:
        par = box(slide, M, 1.25, 5.6, 0.3).paragraphs[0]
        run(par, 'FINAL YEAR PROJECT   ·   UI/UX DESIGN PROPOSAL', MONO, 10, SPOT, spacing=120)
        headline(slide, row['title'], None, M, 1.7, 5.6, 1.3, 66)
        tf = box(slide, M, 3.15, 5.2, 1.0)
        par = tf.paragraphs[0]
        par.line_spacing = 1.08
        run(par, row['body'][0], SERIF, 25, SPOT, italic=True)
        line(slide, M, 4.55, M + 5.2, 4.55, RULE, 0.5)
        par = box(slide, M, 4.72, 5.0, 0.9).paragraphs[0]
        par.line_spacing = 1.2
        run(par, row['body'][1], SANS, 16)
        if team and team.get('members'):
            par = box(slide, M, 5.75, 5.2, 0.9).paragraphs[0]
            par.line_spacing = 1.25
            run(par, '   ·   '.join(m['name'] for m in team['members']), SANS, 12, INK2)
        signal_rule(slide, M, 6.45, 5.2)
        figure(slide, image, 6.55, 1.25, W - M - 6.55, 4.95, 1, CAPTION[1], row['title'] + ' - ' + CAPTION[1])
    elif not image:
        par = box(slide, M, 1.2, 8, 0.3).paragraphs[0]
        run(par, f'{n:02d} — {SECTION[i].upper()}', MONO, 10, SPOT, spacing=120)
        headline(slide, row['title'], EMPHASIS[i], M, 1.6, 10.6, 1.9, 42)
        cols = len(row['body'])
        gap = 0.45
        cw = (W - 2 * M - gap * (cols - 1)) / cols
        for j, text in enumerate(row['body']):
            x = M + j * (cw + gap)
            line(slide, x, 4.0, x + cw, 4.0, INK, 0.75)
            par = box(slide, x, 4.18, cw, 0.8).paragraphs[0]
            run(par, f'{j + 1:02d}', SERIF, 40, SPOT, italic=True)
            par = box(slide, x, 5.1, cw - 0.2, 1.5).paragraphs[0]
            par.line_spacing = 1.18
            run(par, text, SANS, 18)
    else:
        left_w = 4.85
        par = box(slide, M, 1.2, left_w, 0.3).paragraphs[0]
        run(par, f'{n:02d} — {SECTION[i].upper()}', MONO, 10, SPOT, spacing=120)
        headline(slide, row['title'], EMPHASIS[i], M, 1.6, left_w, 2.2, 31)
        points(slide, row['body'], M, 3.95, left_w, 0.92)
        fx = M + left_w + 0.6
        figure(slide, image, fx, 1.25, W - M - fx, 4.95, n, CAPTION[n], row['title'] + ' - ' + CAPTION[n])
    stamp = lambda s: f'{s // 60}:{s % 60:02d}'
    slide.notes_slide.notes_text_frame.text = (
        f"Timing: {row['seconds']} seconds. Segment {stamp(start)}-{stamp(minute)}.\n\n{row['notes']}")

final = OUT / 'PARALLAX-Presentation-Final.pptx'
prs.save(str(final))
(OUT / 'Presenter-Notes.md').write_text(
    '\n'.join(f"## {i + 1}. {r['title']}\n\n{r['seconds']} seconds\n\n{r['notes']}\n" for i, r in enumerate(content)),
    encoding='utf-8', newline=chr(10))
print(json.dumps({'slides': len(content), 'plannedSeconds': minute, 'final': str(final), 'team': bool(team)}))
